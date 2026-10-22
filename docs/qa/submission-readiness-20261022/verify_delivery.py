from pathlib import Path
import hashlib,json,re,unicodedata,zipfile
from lxml import etree as E
from pptx import Presentation
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).resolve().parent
folder=ROOT/'presentations/competition-2026/audited-20261022'
ppt=next(folder.glob('*.pptx'));pdf=next(folder.glob('*.pdf'));video=next(folder.glob('*.mp4'))
pr=Presentation(ppt);reader=PdfReader(pdf)
assert len(pr.slides)==len(reader.pages)==32
norm=lambda s:re.sub(r'\s+','',unicodedata.normalize('NFC',s))
def walk(shapes):
    for sh in shapes:
        if sh.shape_type==6:yield from walk(sh.shapes)
        else:yield sh
rows=[]
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with zipfile.ZipFile(ppt) as z:
    assert not any(re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n) for n in z.namelist())
    assert z.testzip() is None
    for i,s in enumerate(pr.slides):
        x=E.fromstring(z.read(f'ppt/slides/slide{i+1}.xml'))
        assert not x.findall('p:transition',ns) and not x.findall('p:timing',ns)
        for run in x.findall('.//a:r',ns):
            text=run.findtext('a:t',namespaces=ns) or ''
            if not text.strip():continue
            p=run.find('a:rPr',ns);assert p is not None
            assert p.find('a:latin',ns).get('typeface')=='Times New Roman'
            assert p.find('a:ea',ns).get('typeface') in ['Times New Roman','DFKai-SB','標楷體']
            if re.search('[\u3400-\u9fff]',text):assert p.find('a:ea',ns).get('typeface') in ['DFKai-SB','標楷體']
        text=reader.pages[i].extract_text();assert '\0' not in text
        missing=[p.text for sh in walk(s.shapes) if sh.has_text_frame for p in sh.text_frame.paragraphs if p.text.strip() and norm(p.text) not in norm(text)]
        assert not missing,(i+1,missing)
        rows.append({'page':i+1,'pdf_text_complete':True,'fonts_pass':True})
    assert hashlib.sha256(z.read('ppt/media/media1.mp4')).digest()==hashlib.sha256(video.read_bytes()).digest()
for n in [0,31]:
    urls=[str(a.get_object().get('/A',{}).get('/URI','')) for a in reader.pages[n].get('/Annots',[])]
    assert any('github.com' in u for u in urls) and any('github.io' in u for u in urls),(n,urls)
expected={'.pptx':'0edc2dac0fa5bd1df514df931ee668192467754a695d7ce2b07dc7edff3a386a','.pdf':'66f32e392cb00298e27094adb160141e31ceca549a5154a7982ed6ff218fb9a6','.mp4':'da9bde7cf26c4e16f258ee8920f34acb9a1e0b76231fb2f9a608b67b749f38e0'}
files=[]
for f in [ppt,pdf,video]:
    digest=hashlib.sha256(f.read_bytes()).hexdigest();assert digest==expected[f.suffix]
    files.append({'path':f.relative_to(ROOT).as_posix(),'sha256':digest})
r={'passed':True,'pages':rows,'no_notes_animations_transitions':True,'embedded_video_matches':True,'cover_and_end_links':True,'files':files,'native_powerpoint_playback':'not verified','demo':'real UI screenshot montage; not continuous recording'}
(OUT/'delivery-verification.json').write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf8')
print('PASS: 32 slides, fonts, PDF text and links, embedded video, unchanged delivery hashes')
